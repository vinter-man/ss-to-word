from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from ss_to_word.export import ExportOptions, PageText, render_txt, write_docx, write_txt
from ss_to_word.layout import Block


def pages() -> list[PageText]:
    return [
        PageText(Path("IMG_0001.jpg"), [
            Block("heading", "СПОРЕД ФОРМАТА НА ОРГАНИЗАЦИЯ"),
            Block("bullet", "неорганизирано"),
            Block("bullet", "организирано от учреждения, предприятия и училища"),
        ], mean_conf=0.96),
        PageText(Path("IMG_0002.jpg"), [Block("paragraph", "обикновен абзац с текст")], mean_conf=0.9),
        PageText(Path("IMG_0003.jpg"), [], mean_conf=0.0),
        PageText(Path("broken.jpg"), [], error="could not open the image"),
    ]


def test_write_docx_structure(tmp_path: Path):
    out = tmp_path / "result.docx"
    write_docx(pages(), out, ExportOptions(title="Лекция", file_headings=True))
    doc = Document(str(out))
    paras = [(p.style.name, p.text) for p in doc.paragraphs if p.text.strip()]

    assert paras[0] == ("Title", "Лекция")
    assert ("Heading 2", "IMG_0001.jpg") in paras
    assert ("Slide Heading", "СПОРЕД ФОРМАТА НА ОРГАНИЗАЦИЯ") in paras
    assert ("List Bullet", "неорганизирано") in paras
    assert ("Normal", "обикновен абзац с текст") in paras
    assert any("[no text found]" in t for _, t in paras)
    assert any("could not open" in t for _, t in paras)
    # order of images is preserved
    names = [t for s, t in paras if s == "Heading 2"]
    assert names == ["IMG_0001.jpg", "IMG_0002.jpg", "IMG_0003.jpg", "broken.jpg"]
    # Bulgarian proofing language and a Cyrillic-capable font on the base style
    normal = doc.styles["Normal"]
    assert normal.element.rPr.find(qn("w:lang")).get(qn("w:val")) == "bg-BG"
    assert normal.font.name == "Calibri"
    assert doc.core_properties.title == "Лекция"


def test_write_docx_without_file_headings_and_with_page_breaks(tmp_path: Path):
    out = tmp_path / "r.docx"
    write_docx(pages()[:2], out, ExportOptions(title="", file_headings=False, page_breaks=True))
    doc = Document(str(out))
    assert not any(p.style.name == "Heading 2" for p in doc.paragraphs)
    assert not any(p.style.name == "Title" for p in doc.paragraphs)
    xml = doc.element.xml
    assert xml.count('w:type="page"') == 1


def test_render_txt():
    text = render_txt(pages(), ExportOptions(title="Лекция"))
    assert text.startswith("Лекция\n======\n")
    assert "--- IMG_0001.jpg ---" in text
    assert "• неорганизирано\n• организирано от учреждения" in text
    assert "обикновен абзац с текст" in text
    assert "[no text found]" in text
    assert "[could not process this file: could not open the image]" in text
    assert text.index("IMG_0001") < text.index("IMG_0002") < text.index("IMG_0003")


def test_write_txt_utf8(tmp_path: Path):
    out = tmp_path / "nested" / "r.txt"
    write_txt(pages(), out)
    assert out.read_text(encoding="utf-8").count("•") == 2
