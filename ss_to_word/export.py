"""Writers for the recognized text: Word (.docx) and plain text (.txt)."""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from .layout import Block

BULLET = "•"


@dataclass
class PageText:
    """Text recognized from one image."""
    source: Path
    blocks: list[Block] = field(default_factory=list)
    mean_conf: float = 0.0
    rotation: int = 0          # degrees the image had to be turned to be readable
    error: str | None = None   # set when the image could not be processed


@dataclass
class ExportOptions:
    title: str = ""
    file_headings: bool = True          # "IMG_0001.jpg" heading before each image's text
    page_breaks: bool = False           # start every image on a new page
    font_name: str = "Calibri"
    font_size_pt: float = 12
    language: str = "bg-BG"


# ------------------------------------------------------------------------------------- docx
def write_docx(pages: list[PageText], out: Path, opts: ExportOptions | None = None) -> None:
    from docx import Document
    from docx.enum.text import WD_BREAK, WD_LINE_SPACING
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    opts = opts or ExportOptions()
    doc = Document()
    _style_document(doc, opts, Pt, RGBColor, WD_LINE_SPACING, OxmlElement, qn)

    if opts.title:
        doc.add_heading(opts.title, level=0)
        meta = doc.add_paragraph()
        run = meta.add_run(f"Images: {len(pages)} · {datetime.now():%Y-%m-%d %H:%M}")
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    for i, page in enumerate(pages):
        if opts.page_breaks and i > 0:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        if opts.file_headings:
            doc.add_heading(page.source.name, level=2)
        if page.error:
            p = doc.add_paragraph()
            r = p.add_run(f"[could not process this file: {page.error}]")
            r.italic = True
            r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
            continue
        if not page.blocks:
            p = doc.add_paragraph()
            r = p.add_run("[no text found]")
            r.italic = True
            r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
            continue
        for block in page.blocks:
            if block.kind == "heading":
                p = doc.add_paragraph(style="Slide Heading")
                p.add_run(block.text)
            elif block.kind == "bullet":
                doc.add_paragraph(block.text, style="List Bullet")
            else:
                doc.add_paragraph(block.text)

    doc.core_properties.title = opts.title or "OCR"
    doc.core_properties.language = opts.language
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))


def _style_document(doc, opts: ExportOptions, Pt, RGBColor, WD_LINE_SPACING, OxmlElement, qn) -> None:
    from docx.enum.style import WD_STYLE_TYPE

    normal = doc.styles["Normal"]
    normal.font.name = opts.font_name
    normal.font.size = Pt(opts.font_size_pt)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.15
    _set_lang(normal, opts.language, OxmlElement, qn)
    _set_east_asian_font(normal, opts.font_name, qn)

    bullets = doc.styles["List Bullet"]
    bullets.font.name = opts.font_name
    bullets.font.size = Pt(opts.font_size_pt)
    bullets.paragraph_format.space_after = Pt(3)

    file_heading = doc.styles["Heading 2"]
    file_heading.font.name = opts.font_name
    file_heading.font.size = Pt(10)
    file_heading.font.bold = False
    file_heading.font.italic = True
    file_heading.font.color.rgb = RGBColor(0x8A, 0x8A, 0x8A)
    file_heading.paragraph_format.space_before = Pt(18)
    file_heading.paragraph_format.space_after = Pt(4)

    slide = doc.styles.add_style("Slide Heading", WD_STYLE_TYPE.PARAGRAPH)
    slide.base_style = normal
    slide.font.name = opts.font_name
    slide.font.size = Pt(opts.font_size_pt + 2)
    slide.font.bold = True
    slide.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    slide.paragraph_format.space_before = Pt(10)
    slide.paragraph_format.space_after = Pt(6)
    slide.paragraph_format.keep_with_next = True

    title = doc.styles["Title"]
    title.font.name = opts.font_name


def _set_lang(style, lang: str, OxmlElement, qn) -> None:
    rpr = style.element.get_or_add_rPr()
    el = rpr.find(qn("w:lang"))
    if el is None:
        el = OxmlElement("w:lang")
        rpr.append(el)
    el.set(qn("w:val"), lang)


def _set_east_asian_font(style, font: str, qn) -> None:
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is not None:
        fonts.set(qn("w:eastAsia"), font)
        fonts.set(qn("w:cs"), font)


# -------------------------------------------------------------------------------------- txt
def render_txt(pages: list[PageText], opts: ExportOptions | None = None) -> str:
    opts = opts or ExportOptions()
    out: list[str] = []
    if opts.title:
        out += [opts.title, "=" * min(60, len(opts.title)), ""]
    for page in pages:
        if opts.file_headings:
            out += [f"--- {page.source.name} ---", ""]
        if page.error:
            out += [f"[could not process this file: {page.error}]", ""]
            continue
        if not page.blocks:
            out += ["[no text found]", ""]
            continue
        for block in page.blocks:
            if block.kind == "bullet":
                out += [f"{BULLET} {block.text}"]
            else:
                if out and out[-1] != "":
                    out.append("")          # keep a blank line after a bullet list
                out += [block.text, ""]
        if out and out[-1] != "":
            out.append("")
        out.append("")
    return "\n".join(out).rstrip() + "\n"


def write_txt(pages: list[PageText], out: Path, opts: ExportOptions | None = None) -> None:
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(render_txt(pages, opts), encoding="utf-8")
